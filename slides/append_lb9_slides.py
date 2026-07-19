"""Append the newer analyses to exp20_lb9_generalization.pptx (keep the 11 confirmed slides):
  A. effective-OE eval (|self-LFC|>=0.25, n=27)   B. lineage clustering.
Then move 'Next steps' to the very end. Operates on the OUTPUT pptx only.
"""
import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
DECK = os.path.join(os.path.dirname(__file__), 'exp20_lb9_generalization.pptx')
FIG = '/home/sgoudarzi/scCausalRL/experiments/exp20_load_bearing/results/figures'
WHITE = RGBColor(0xFF, 0xFF, 0xFF); BLACK = RGBColor(0, 0, 0); DARK = RGBColor(0x1A, 0x1A, 0x2E)
DARKBLUE = RGBColor(0x0B, 0x3D, 0x91); GRAY = RGBColor(0x55, 0x55, 0x55); GREEN = RGBColor(0x14, 0x60, 0x14)
prs = Presentation(DECK)


def blank():
    s = prs.slides.add_slide(prs.slide_layouts[6]); s.background.fill.solid(); s.background.fill.fore_color.rgb = WHITE; return s


def text(s, l, t, w, h, txt, *, size=18, bold=False, color=BLACK):
    tb = s.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h)); tf = tb.text_frame; tf.word_wrap = True
    p = tf.paragraphs[0]; p.text = txt; p.font.size = Pt(size); p.font.bold = bold; p.font.color.rgb = color; p.font.name = 'Calibri'; return tb


def title_bar(s, txt):
    sh = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.85))
    sh.fill.solid(); sh.fill.fore_color.rgb = DARK; sh.line.fill.background(); sh.shadow.inherit = False
    text(s, 0.5, 0.15, 12.3, 0.6, txt, size=23, bold=True, color=WHITE)


def box(s, l, t, w, h, fill, line):
    sh = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(l), Inches(t), Inches(w), Inches(h))
    sh.fill.solid(); sh.fill.fore_color.rgb = fill; sh.line.color.rgb = line; sh.line.width = Pt(1.5); sh.shadow.inherit = False


def figim(s, l, t, path, w):
    if os.path.exists(path): s.shapes.add_picture(path, Inches(l), Inches(t), width=Inches(w))


def foot(s, txt): text(s, 0.5, 7.05, 12.3, 0.35, txt, size=11, color=GRAY)


# A. effective-OE eval (n=27)
s = blank(); title_bar(s, 'Eval on effective perturbations  (|self-LFC| ≥ 0.25,  n = 27)')
figim(s, 0.35, 1.05, f'{FIG}/best_model_lfc025.png', 12.6)
box(s, 0.4, 5.95, 12.55, 1.0, RGBColor(0xEC, 0xF7, 0xEC), RGBColor(0x27, 0xAE, 0x60))
text(s, 0.6, 6.05, 12.15, 0.85,
     '27/49 held-out TFs have effective OE. Restricting to them WIDENS the copy-beating margin: top '
     'model allDEG 0.41 vs ESM-kNN 0.28 (all-49: 0.35 vs 0.31). Weak-OE TFs (near-null truth) dilute '
     'the aggregate — an eval fix, not a training change.', size=13, bold=True, color=GREEN)
foot(s, 'Held-out singles, 4000 control cells, drift-removed, ±1 SEM. Filter is on the perturbation (OE efficiency), not the DEG threshold (0.1).')

# B. lineage clustering
s = blank(); title_bar(s, 'Where the model works — held-out TFs by lineage  (n = 46 annotated)')
figim(s, 0.35, 1.05, f'{FIG}/tf_fate_cluster.png', 12.6)
box(s, 0.4, 5.95, 12.55, 1.0, RGBColor(0xEC, 0xF1, 0xFB), RGBColor(0x21, 0x71, 0xB5))
text(s, 0.6, 6.05, 12.15, 0.85,
     'Prediction varies by the lineage the TF drives: Epithelial/Endoderm best (allDEG 0.77), '
     'Mesenchyme/Muscle worst (0.17), Neural mid (0.36). Fate-COMMITMENT correlates +0.12 with allDEG — '
     'committed TFs are slightly EASIER, not harder (small n/lineage; wide CIs).', size=13, bold=True, color=DARKBLUE)
foot(s, 'main_fate / fate-specificity from the Cao mouse-organogenesis developmental annotation; allDEG from the top model (mmd1_deg2).')

# move original 'Next steps' (11th slide, index 10) to the end
sl = prs.slides._sldIdLst; ids = list(sl); nxt = ids[10]; sl.remove(nxt); sl.append(nxt)
prs.save(DECK)
print('saved', DECK, '|', len(prs.slides._sldIdLst), 'slides')
