"""Render a .pptx to per-slide PNGs (matplotlib) by reading shape geometry — no LibreOffice needed.
Approximate (font metrics/wrap are estimated) but faithful enough to catch overcrowding / overlap / missing figures.
Usage: python slides/render_pptx_preview.py <deck.pptx> <out_dir>
"""
import sys, os, io, textwrap
import matplotlib; matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib.patches import Rectangle, FancyBboxPatch
from PIL import Image
from pptx import Presentation
from pptx.util import Emu
from pptx.enum.shapes import MSO_SHAPE_TYPE

EMU = 914400.0
deck = sys.argv[1]; outdir = sys.argv[2]; os.makedirs(outdir, exist_ok=True)
prs = Presentation(deck); SW = prs.slide_width / EMU; SH = prs.slide_height / EMU


def hexof(color, default=None):
    try:
        return '#' + str(color.rgb)
    except Exception:
        return default


def _fontcol(font):
    # Font color lives at font.color.rgb (NOT font.rgb) — the old hexof(font) always failed -> black.
    try:
        return '#' + str(font.color.rgb)
    except Exception:
        return None


def para_props(para):
    sz = None; bold = bool(para.font.bold); col = _fontcol(para.font)
    if para.font.size is not None: sz = para.font.size.pt
    for r in para.runs:
        if sz is None and r.font.size is not None: sz = r.font.size.pt
        if r.font.bold: bold = True
        if col is None: col = _fontcol(r.font)
    return sz or 14, bold, (col or '#000000')


for si, slide in enumerate(prs.slides, 1):
    fig = plt.figure(figsize=(SW, SH)); ax = fig.add_axes([0, 0, 1, 1])
    ax.set_xlim(0, SW); ax.set_ylim(SH, 0); ax.axis('off')
    ax.add_patch(Rectangle((0, 0), SW, SH, facecolor='white', zorder=0))
    for sh in slide.shapes:
        try:
            l, t, w, h = sh.left / EMU, sh.top / EMU, (sh.width or 0) / EMU, (sh.height or 0) / EMU
        except Exception:
            continue
        # picture
        if sh.shape_type == MSO_SHAPE_TYPE.PICTURE:
            try:
                im = Image.open(io.BytesIO(sh.image.blob)); ax.imshow(im, extent=(l, l + w, t + h, t), zorder=3, aspect='auto')
            except Exception:
                ax.add_patch(Rectangle((l, t), w, h, facecolor='#dddddd', edgecolor='#999', zorder=3))
            continue
        # table
        if sh.has_table:
            tb = sh.table; nr = len(tb.rows); nc = len(tb.columns)
            cw = [c.width / EMU for c in tb.columns]; rh = h / nr
            y = t
            for i in range(nr):
                x = l
                for j in range(nc):
                    cell = tb.cell(i, j); fc = hexof(cell.fill.fore_color, '#ffffff') if cell.fill.type == 1 else '#ffffff'
                    ax.add_patch(Rectangle((x, y), cw[j], rh, facecolor=fc, edgecolor='#cccccc', lw=0.4, zorder=2))
                    p = cell.text_frame.paragraphs[0]; sz, bold, colr = para_props(p)
                    ax.text(x + 0.04, y + rh / 2, cell.text, fontsize=min(sz, 11) * 0.9, color=colr,
                            va='center', ha='left', zorder=4, wrap=True)
                    x += cw[j]
                y += rh
            continue
        # autoshape fill (rounded rect etc.)
        is_text = sh.has_text_frame
        fc = None
        if sh.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE:
            try:
                if sh.fill.type == 1: fc = hexof(sh.fill.fore_color, None)
            except Exception:
                fc = None
            if fc: ax.add_patch(FancyBboxPatch((l + 0.02, t + 0.02), max(w - 0.04, 0.01), max(h - 0.04, 0.01),
                                boxstyle='round,pad=0.02', facecolor=fc, edgecolor='none', zorder=1))
        # text
        if is_text and sh.text_frame.text.strip():
            y = t + 0.06
            for p in sh.text_frame.paragraphs:
                if not p.text.strip():
                    y += 0.14; continue
                sz, bold, colr = para_props(p)
                # chars-per-line ≈ box_width / char_width; bold chars are wider. (old 155 over-estimated
                # -> text ran past the box; ~120/105 matches PowerPoint word-wrap far more closely.)
                cpl = max(int(w * (105 if bold else 120) / max(sz, 8)), 6)
                wrapped = textwrap.fill(p.text, width=cpl)
                nlines = wrapped.count('\n') + 1
                ha = {1: 'left', 2: 'center', 3: 'right'}.get(int(p.alignment) if p.alignment is not None else 1, 'left')
                xx = l + w / 2 if ha == 'center' else (l + w - 0.05 if ha == 'right' else l + 0.08)
                ax.text(xx, y, wrapped, fontsize=sz * 0.92, color=colr, va='top', ha=ha,
                        weight='bold' if bold else 'normal', zorder=5)
                y += nlines * sz / 72.0 * 1.25 + 0.04
    out = os.path.join(outdir, f'slide{si:02d}.png')
    fig.savefig(out, dpi=110, facecolor='white'); plt.close(fig)
    print('wrote', out)
print('done', len(prs.slides._sldIdLst), 'slides')
